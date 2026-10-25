# -*- coding: utf-8 -*-
"""
generar_plantilla_word.py
=========================
Crea assets/plantillas/plantilla_unicartagena.docx: la plantilla de estilos (reference-doc)
que Pandoc usa para dar formato **APA 7.ª edición (trabajo de estudiante)** a los dos estudios.

  * Carta (21,59 × 27,94 cm), márgenes de 2,54 cm (1 pulgada)
  * Times New Roman 12, texto en negro, interlineado doble, alineado a la izquierda,
    sangría de primera línea de 1,27 cm, sin espacio extra entre párrafos
  * Títulos APA: nivel 1 centrado y negrita; nivel 2 a la izquierda y negrita;
    nivel 3 a la izquierda, negrita y cursiva. Todo en negro.
  * Número de página arriba a la derecha (encabezado)
  * Estilos de portada (Portada) y de referencias (Referencia, con sangría francesa)

Uso: python scripts/generar_plantilla_word.py
"""
from __future__ import annotations

import os
import subprocess

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = os.path.join(RAIZ, "assets", "plantillas", "plantilla_unicartagena.docx")

NEGRO = RGBColor(0, 0, 0)
FUENTE = "Times New Roman"
SANGRIA = Cm(1.27)


def fijar_fuente(style, tam=12, negrita=None, cursiva=None):
    f = style.font
    f.name = FUENTE
    f.size = Pt(tam)
    f.color.rgb = NEGRO
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(att), FUENTE)
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(att)) is not None:
            del rfonts.attrib[qn(att)]
    if negrita is not None:
        f.bold = negrita
    if cursiva is not None:
        f.italic = cursiva


def parrafo(style, alineacion=WD_ALIGN_PARAGRAPH.LEFT, sangria=None, doble=True, colgante=None):
    pf = style.paragraph_format
    pf.alignment = alineacion
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE if doble else WD_LINE_SPACING.SINGLE
    pf.first_line_indent = sangria
    if colgante is not None:
        pf.left_indent = colgante
        pf.first_line_indent = -colgante


def estilo(doc, nombre):
    try:
        return doc.styles[nombre]
    except KeyError:
        st = doc.styles.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
        return st


def quitar_borde(st):
    pPr = st.element.get_or_add_pPr()
    for b in pPr.findall(qn("w:pBdr")):
        pPr.remove(b)


def campo_pagina(parrafo_):
    run = parrafo_.add_run()
    run.font.name = FUENTE
    run.font.size = Pt(12)
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


def main():
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    base = SALIDA + ".base.docx"
    with open(base, "wb") as f:
        subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], stdout=f, check=True)
    doc = Document(base)

    # Página: carta, 1 pulgada, número de página arriba a la derecha
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.59), Cm(27.94)
        for lado in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(s, lado, Cm(2.54))
        s.header_distance = s.footer_distance = Cm(1.27)
        for p in s.footer.paragraphs:
            p.text = ""
        enc = s.header.paragraphs[0] if s.header.paragraphs else s.header.add_paragraph()
        enc.text = ""
        enc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        campo_pagina(enc)

    # Texto
    for nombre in ("Normal", "Body Text", "First Paragraph"):
        st = estilo(doc, nombre)
        fijar_fuente(st)
        parrafo(st, sangria=SANGRIA if nombre != "Normal" else None)
    st = estilo(doc, "Compact")  # celdas de tablas y listas
    fijar_fuente(st)
    parrafo(st, doble=False, sangria=Cm(0))
    for nombre in ("Block Text", "Footnote Text", "Definition", "Definition Term"):
        st = estilo(doc, nombre)
        fijar_fuente(st)

    # Títulos APA 7
    niveles = {
        1: dict(al=WD_ALIGN_PARAGRAPH.CENTER, neg=True, cur=False),
        2: dict(al=WD_ALIGN_PARAGRAPH.LEFT, neg=True, cur=False),
        3: dict(al=WD_ALIGN_PARAGRAPH.LEFT, neg=True, cur=True),
        4: dict(al=WD_ALIGN_PARAGRAPH.LEFT, neg=True, cur=False),
        5: dict(al=WD_ALIGN_PARAGRAPH.LEFT, neg=True, cur=True),
    }
    for n, cfg in niveles.items():
        st = estilo(doc, f"Heading {n}")
        fijar_fuente(st, negrita=cfg["neg"], cursiva=cfg["cur"])
        parrafo(st, alineacion=cfg["al"], sangria=Cm(0))
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.keep_with_next = True
        quitar_borde(st)
    for nombre in ("Title", "Subtitle", "TOC Heading"):
        st = estilo(doc, nombre)
        fijar_fuente(st, negrita=True, cursiva=False)
        parrafo(st, alineacion=WD_ALIGN_PARAGRAPH.CENTER)
        quitar_borde(st)

    # Leyendas (el compilador las divide en "Tabla N" en negrita + título en cursiva)
    for nombre in ("Table Caption", "Image Caption"):
        st = estilo(doc, nombre)
        fijar_fuente(st, negrita=False, cursiva=True)
        parrafo(st, sangria=Cm(0))
        st.paragraph_format.keep_with_next = True
    st = estilo(doc, "Captioned Figure")
    parrafo(st, alineacion=WD_ALIGN_PARAGRAPH.CENTER, doble=False)

    # Portada (APA estudiante: todo centrado, doble espacio, 12 pt)
    st = estilo(doc, "Portada")
    fijar_fuente(st)
    parrafo(st, alineacion=WD_ALIGN_PARAGRAPH.CENTER)

    # Referencias: sangría francesa de 1,27 cm, doble espacio
    st = estilo(doc, "Referencia")
    fijar_fuente(st)
    parrafo(st, colgante=SANGRIA)

    # Notas de tablas y figuras
    st = estilo(doc, "Nota")
    fijar_fuente(st)
    parrafo(st, sangria=Cm(0))
    st.paragraph_format.space_after = Pt(12)

    # Hipervínculos en negro (APA: URL como texto normal)
    try:
        st = doc.styles["Hyperlink"]
        st.font.color.rgb = NEGRO
        st.font.underline = False
    except KeyError:
        pass

    doc.save(SALIDA)
    os.remove(base)
    print("OK ·", os.path.relpath(SALIDA, RAIZ))


if __name__ == "__main__":
    main()
