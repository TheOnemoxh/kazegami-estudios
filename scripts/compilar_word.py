# -*- coding: utf-8 -*-
"""
compilar_word.py
================
Une los .md de cada estudio con Pandoc, aplica la plantilla APA 7 y deja listos:

    entregables_word/Taller_1_Estudio_de_Mercado_Kazegami.docx
    entregables_word/Taller_2_Estudio_Tecnico_Kazegami.docx

Después de Pandoc, python-docx aplica las normas APA 7.ª edición:
  * tabla de contenido después de la portada (se actualiza al abrir en Word);
  * tablas APA: "Tabla N" en negrita y, debajo, el título en cursiva; solo líneas horizontales
    (arriba, bajo el encabezado y al final), sin color ni sombreado, texto en negro;
  * figuras APA: "Figura N" en negrita y el título en cursiva ARRIBA de la imagen;
  * notas "Nota." alineadas a la izquierda y sin sangría, debajo de tablas y figuras.
Al final verifica que cada .docx exista, pese lo esperado y contenga todas las tablas e imágenes.

Uso:
    python scripts/compilar_word.py            # compila los dos
    python scripts/compilar_word.py --pdf      # además exporta PDF con LibreOffice (opcional)
"""
from __future__ import annotations

import copy
import glob
import os
import re
import shutil
import subprocess
import sys
import zipfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PLANTILLA = os.path.join(RAIZ, "assets", "plantillas", "plantilla_unicartagena.docx")
SALIDA = os.path.join(RAIZ, "entregables_word")

DOCUMENTOS = [
    ("01_estudio_mercado", "Taller_1_Estudio_de_Mercado_Kazegami.docx", "Taller No. 1 - Estudio de Mercados - Kazegami Restaurante"),
    ("02_estudio_tecnico", "Taller_2_Estudio_Tecnico_Kazegami.docx", "Taller No. 2 - Estudio Técnico - Kazegami Restaurante"),
]

NEGRO = RGBColor(0, 0, 0)
FUENTE = "Times New Roman"


# ---------------------------------------------------------------------------
def pandoc(carpeta: str, salida: str, titulo: str):
    archivos = sorted(glob.glob(os.path.join(RAIZ, carpeta, "*.md")))
    if not archivos:
        raise SystemExit(f"No hay .md en {carpeta}")
    cmd = ["pandoc", *archivos, "-f", "markdown+pipe_tables+raw_attribute+fenced_divs-tex_math_dollars-tex_math_single_backslash",
           "-t", "docx", "--reference-doc", PLANTILLA, "--resource-path", RAIZ,
           "--metadata", f"title-meta={titulo}", "-o", salida]
    subprocess.run(cmd, check=True, cwd=RAIZ)
    return archivos


def _borde(nombre, val="single", sz="8"):
    e = OxmlElement(f"w:{nombre}")
    e.set(qn("w:val"), val)
    if val != "nil":
        e.set(qn("w:sz"), sz)
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
    return e


def _run_fmt(r, tam=None, negrita=None, cursiva=None):
    r.font.name = FUENTE
    r._r.get_or_add_rPr()
    rfonts = r._r.rPr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        r._r.rPr.insert(0, rfonts)
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(att), FUENTE)
    r.font.color.rgb = NEGRO
    if tam:
        r.font.size = tam
    if negrita is not None:
        r.font.bold = negrita
    if cursiva is not None:
        r.font.italic = cursiva


def formatear_tablas_apa(doc):
    """APA 7: solo bordes horizontales (superior, bajo encabezados, inferior); sin verticales ni sombreado."""
    for t in doc.tables:
        tblPr = t._tbl.tblPr
        for tag in ("w:tblStyle", "w:tblBorders", "w:tblW", "w:tblCellMar", "w:shd"):
            for e in tblPr.findall(qn(tag)):
                tblPr.remove(e)
        borders = OxmlElement("w:tblBorders")
        borders.append(_borde("top"))
        borders.append(_borde("left", "nil"))
        borders.append(_borde("bottom"))
        borders.append(_borde("right", "nil"))
        borders.append(_borde("insideH", "nil"))
        borders.append(_borde("insideV", "nil"))
        tblPr.append(borders)
        tblW = OxmlElement("w:tblW")
        tblW.set(qn("w:w"), "5000")
        tblW.set(qn("w:type"), "pct")
        tblPr.append(tblW)
        mar = OxmlElement("w:tblCellMar")
        for lado, v in (("top", "20"), ("bottom", "20"), ("left", "60"), ("right", "60")):
            e = OxmlElement(f"w:{lado}")
            e.set(qn("w:w"), v)
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        tblPr.append(mar)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        # espacio después de la tabla cuando no sigue una nota
        sig = t._tbl.getnext()
        if sig is not None and sig.tag == qn("w:p"):
            sp = Paragraph(sig, t._parent)
            if not sp.text.strip().startswith("Nota."):
                sp.paragraph_format.space_before = Pt(12)

        ncol = len(t.columns)
        tam = Pt(8) if ncol >= 7 else Pt(9) if ncol >= 5 else Pt(10)
        for i, fila in enumerate(t.rows):
            trPr = fila._tr.get_or_add_trPr()
            if i == 0:
                h = OxmlElement("w:tblHeader")
                h.set(qn("w:val"), "true")
                trPr.append(h)
            trPr.append(OxmlElement("w:cantSplit"))
            for celda in fila.cells:
                tcPr = celda._tc.get_or_add_tcPr()
                for e in tcPr.findall(qn("w:shd")) + tcPr.findall(qn("w:tcBorders")):
                    tcPr.remove(e)
                if i == 0:
                    tb = OxmlElement("w:tcBorders")
                    tb.append(_borde("top"))
                    tb.append(_borde("bottom"))
                    tcPr.append(tb)
                for p in celda.paragraphs:
                    pf = p.paragraph_format
                    pf.space_after = Pt(0)
                    pf.space_before = Pt(0)
                    pf.line_spacing = 1.0
                    pf.first_line_indent = Pt(0)
                    pf.left_indent = Pt(0)
                    if i == 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        _run_fmt(r, tam=tam, negrita=False if i == 0 else None)


def _nuevo_parrafo_antes(ref_el, estilo, partes, padre):
    """Crea un párrafo antes de ref_el. partes = [(texto, negrita, cursiva)]."""
    p = OxmlElement("w:p")
    ref_el.addprevious(p)
    par = Paragraph(p, padre)
    par.style = estilo
    for texto, neg, cur in partes:
        r = par.add_run(texto)
        _run_fmt(r, tam=Pt(12), negrita=neg, cursiva=cur)
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.first_line_indent = Pt(0)
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return par


def leyendas_apa(doc):
    """Tabla N (negrita) / Título (cursiva) arriba de tablas y figuras."""
    patron = re.compile(r"^(Tabla|Figura)\s+(\d+)\.\s*(.*)$", re.S)
    est_tabla = doc.styles["Table Caption"]
    est_fig = doc.styles["Image Caption"]
    for p in list(doc.paragraphs):
        nombre = p.style.name if p.style is not None else ""
        if nombre not in ("Table Caption", "Image Caption"):
            continue
        m = patron.match(p.text.strip())
        if not m:
            continue
        tipo, num, titulo = m.groups()
        if tipo == "Tabla":
            ancla = p._p
            _nuevo_parrafo_antes(ancla, est_tabla, [(f"Tabla {num}", True, False)], p._parent)
            _nuevo_parrafo_antes(ancla, est_tabla, [(titulo, False, True)], p._parent)
            ancla.getparent().remove(ancla)
        else:
            # La imagen es el párrafo anterior; la leyenda debe ir ARRIBA de la imagen
            img = p._p.getprevious()
            while img is not None and not img.xpath(".//w:drawing"):
                img = img.getprevious()
            ancla = img if img is not None else p._p
            _nuevo_parrafo_antes(ancla, est_fig, [(f"Figura {num}", True, False)], p._parent)
            _nuevo_parrafo_antes(ancla, est_fig, [(titulo, False, True)], p._parent)
            p._p.getparent().remove(p._p)


def formatear_parrafos(doc):
    est_nota = doc.styles["Nota"]
    for p in doc.paragraphs:
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = 1.0
        if p.text.strip().startswith("Nota."):
            p.style = est_nota
            p.paragraph_format.first_line_indent = Pt(0)
            for r in p.runs:
                _run_fmt(r)
        for r in p.runs:
            if r.font.color is not None and r.font.color.rgb not in (None, NEGRO):
                r.font.color.rgb = NEGRO


def insertar_indice(doc):
    """Inserta 'Contenido' + campo TOC justo después del primer salto de página (fin de la portada)."""
    destino = None
    for p in doc.paragraphs:
        if p._p.xpath('.//w:br[@w:type="page"]'):
            destino = p
            break
    if destino is None:
        return
    titulo = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    ps = OxmlElement("w:pStyle")
    ps.set(qn("w:val"), "TOCHeading")
    pPr.append(ps)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    pPr.append(jc)
    titulo.append(pPr)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Contenido"
    r.append(t)
    titulo.append(r)

    campo = OxmlElement("w:p")
    for tipo, instr in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), ("text", None), ("end", None)):
        run = OxmlElement("w:r")
        if tipo == "text":
            tt = OxmlElement("w:t")
            tt.text = "Haga clic derecho → Actualizar campo para ver la tabla de contenido."
            run.append(tt)
        elif tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            if tipo == "begin":
                fc.set(qn("w:dirty"), "true")
            run.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = instr
            run.append(it)
        campo.append(run)

    salto = OxmlElement("w:p")
    rr = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    rr.append(br)
    salto.append(rr)

    destino._p.addnext(salto)
    destino._p.addnext(campo)
    destino._p.addnext(titulo)

    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)


def verificar(ruta, archivos_md):
    md = "\n".join(open(a, encoding="utf-8").read() for a in archivos_md)
    tablas_esp = len(re.findall(r"^Table: ", md, flags=re.M))
    refs_img = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", md)
    doc = Document(ruta)
    with zipfile.ZipFile(ruta) as z:
        imgs = [n for n in z.namelist() if n.startswith("word/media/")]
    peso = os.path.getsize(ruta) / 1024
    tablas = len(doc.tables)
    ley_t = sum(1 for p in doc.paragraphs if re.fullmatch(r"Tabla \d+", p.text.strip()))
    ley_f = sum(1 for p in doc.paragraphs if re.fullmatch(r"Figura \d+", p.text.strip()))
    ok = (peso > 100 and tablas >= tablas_esp and len(imgs) >= len(set(refs_img))
          and ley_t == tablas_esp and ley_f == len(refs_img))
    print(f"  · {os.path.basename(ruta)}: {peso:,.0f} KB · tablas {tablas}/{tablas_esp} (leyendas APA {ley_t}) · "
          f"figuras {ley_f}/{len(refs_img)} · imágenes únicas {len(imgs)} · {'OK' if ok else 'REVISAR'}")
    if "Estudio_Tecnico" in ruta:
        p6 = sum(1 for t in doc.tables if "Identificación del elemento" in " ".join(c.text for c in t.rows[0].cells))
        print(f"    tablas con formato Punto 6 (columna 'Identificación del elemento'): {p6}")
        ok = ok and p6 >= 3
    return ok


def main():
    if not shutil.which("pandoc"):
        raise SystemExit("Falta Pandoc: https://pandoc.org/installing.html")
    subprocess.run([sys.executable, os.path.join(RAIZ, "scripts", "generar_plantilla_word.py")], check=True)
    os.makedirs(SALIDA, exist_ok=True)
    todo_ok = True
    for carpeta, nombre, titulo in DOCUMENTOS:
        ruta = os.path.join(SALIDA, nombre)
        archivos = pandoc(carpeta, ruta, titulo)
        doc = Document(ruta)
        leyendas_apa(doc)
        insertar_indice(doc)
        formatear_tablas_apa(doc)
        formatear_parrafos(doc)
        doc.core_properties.title = titulo
        doc.core_properties.author = "Angulo, Molina, Cañas, Vergara y Rodriguez"
        doc.save(ruta)
        print(f"Compilado {nombre} ({len(archivos)} archivos .md)")
        todo_ok &= verificar(ruta, archivos)
        if "--pdf" in sys.argv and shutil.which("soffice"):
            subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", SALIDA, ruta],
                           check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print("VERIFICACIÓN:", "TODO OK" if todo_ok else "HAY ALGO QUE REVISAR")


if __name__ == "__main__":
    main()
